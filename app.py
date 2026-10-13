import io
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.utils import column_index_from_string

st.set_page_config(
    page_title="MHJ Invoice list",
    page_icon="📊",
    layout="wide",
)

st.title("📊 MHJ Invoice list")
st.markdown(
    "Upload **File A (Source/Lookup)** and **File B (Target Invoice)** to cross-reference PO numbers "
    "and map values while preserving all formatting, styles, and formulas."
)

# ----------------- 1. File Uploaders -----------------
col_file_a, col_file_b = st.columns(2)

with col_file_a:
    file_a = st.file_uploader(
        "Upload File A (Lookup Source - e.g., MHJ SPO)",
        type=["xlsx", "xlsm"],
        key="file_a",
    )

with col_file_b:
    file_b = st.file_uploader(
        "Upload File B (Target to Update - e.g., MJ Invoice)",
        type=["xlsx", "xlsm"],
        key="file_b",
    )

# ----------------- 2. Settings & Coordinate Configuration -----------------
st.markdown("---")
st.subheader("⚙️ Coordinate & Mapping Settings")

# Inspect sheet names dynamically if files are uploaded
sheet_names_a = ["Tag"]
sheet_names_b = ["Invoice"]

if file_a:
    try:
        temp_wb_a = openpyxl.load_workbook(file_a, read_only=True)
        sheet_names_a = temp_wb_a.sheetnames
        file_a.seek(0)
    except Exception:
        pass

if file_b:
    try:
        temp_wb_b = openpyxl.load_workbook(file_b, read_only=True)
        sheet_names_b = temp_wb_b.sheetnames
        file_b.seek(0)
    except Exception:
        pass

config_col1, config_col2 = st.columns(2)

with config_col1:
    st.markdown("#### File A Settings (Source)")
    default_a_idx = (
        sheet_names_a.index("Tag") if "Tag" in sheet_names_a else 0
    )
    selected_sheet_a = st.selectbox(
        "File A Sheet", sheet_names_a, index=default_a_idx
    )
    po_col_a = (
        st.text_input(
            "PO Column in File A",
            value="B",
            help="Column containing PO numbers to match against",
        )
        .strip()
        .upper()
    )
    start_row_a = st.number_input(
        "Data Start Row in File A", min_value=1, value=2, step=1
    )

with config_col2:
    st.markdown("#### File B Settings (Target)")
    selected_sheet_b = st.selectbox("File B Sheet", sheet_names_b, index=0)
    po_col_b = (
        st.text_input(
            "PO Column in File B",
            value="I",
            help="Column containing PO numbers to search for",
        )
        .strip()
        .upper()
    )
    start_row_b = st.number_input(
        "Data Start Row in File B", min_value=1, value=25, step=1
    )

# ----------------- 3. Column Mapping Table -----------------
st.markdown("#### 🔄 Column Mappings (File A ➔ File B)")
st.caption(
    "Define which column in File A maps to which column in File B. You can modify, add, or remove rows below:"
)

# Default preset mappings based on the original request
default_mappings = [
    {
        "File A Column": "E",
        "File B Column": "S",
        "Field Description": "Colour Code",
    },
    {
        "File A Column": "F",
        "File B Column": "T",
        "Field Description": "Size Code",
    },
    {
        "File A Column": "G",
        "File B Column": "L",
        "Field Description": "Barcode",
    },
    {
        "File A Column": "H",
        "File B Column": "M",
        "Field Description": "Item Description",
    },
    {
        "File A Column": "I",
        "File B Column": "K",
        "Field Description": "Style Code",
    },
    {
        "File A Column": "O",
        "File B Column": "R",
        "Field Description": "Product Code",
    },
]

mapping_df = pd.DataFrame(default_mappings)
edited_mapping_df = st.data_editor(
    mapping_df,
    num_rows="dynamic",
    width="stretch",
    column_config={
        "File A Column": st.column_config.TextColumn(
            "File A Source Col", required=True
        ),
        "File B Column": st.column_config.TextColumn(
            "File B Target Col", required=True
        ),
        "Field Description": st.column_config.TextColumn("Description / Label"),
    },
)

# ----------------- 4. Processing Engine -----------------
st.markdown("---")
process_btn = st.button("🚀 Process & Update File B", type="primary")

if process_btn:
    if not file_a or not file_b:
        st.error("⚠️ Please upload both File A and File B to proceed.")
    else:
        with st.spinner("Processing files and applying mappings..."):
            try:
                # 1. Parse and validate column mappings
                col_mappings = []
                for _, row in edited_mapping_df.iterrows():
                    src = str(row["File A Column"]).strip().upper()
                    tgt = str(row["File B Column"]).strip().upper()
                    if src and tgt:
                        src_idx = column_index_from_string(src)
                        tgt_idx = column_index_from_string(tgt)
                        col_mappings.append(
                            (
                                src_idx,
                                tgt_idx,
                                src,
                                tgt,
                                row.get("Field Description", ""),
                            )
                        )

                po_idx_a = column_index_from_string(po_col_a)
                po_idx_b = column_index_from_string(po_col_b)

                # 2. Read File A (Source) with evaluated formula values
                file_a.seek(0)
                wb_a = openpyxl.load_workbook(file_a, data_only=True)
                ws_a = (
                    wb_a[selected_sheet_a]
                    if selected_sheet_a in wb_a.sheetnames
                    else wb_a.active
                )

                po_lookup = {}
                for r in range(start_row_a, ws_a.max_row + 1):
                    po_val = ws_a.cell(row=r, column=po_idx_a).value
                    if po_val is not None:
                        po_key = str(po_val).strip()
                        if po_key and po_key not in po_lookup:
                            # Store target values for each mapped source column
                            po_lookup[po_key] = {
                                src_idx: ws_a.cell(
                                    row=r, column=src_idx
                                ).value
                                for src_idx, _, _, _, _ in col_mappings
                            }

                # 3. Read File B (Target) preserving formatting and formulas
                file_b.seek(0)
                wb_b = openpyxl.load_workbook(file_b, data_only=False)
                ws_b = (
                    wb_b[selected_sheet_b]
                    if selected_sheet_b in wb_b.sheetnames
                    else wb_b.active
                )

                updated_count = 0
                missing_records = []

                for r in range(start_row_b, ws_b.max_row + 1):
                    po_val = ws_b.cell(row=r, column=po_idx_b).value
                    if po_val is None or not str(po_val).strip():
                        continue

                    po_key = str(po_val).strip()
                    if po_key in po_lookup:
                        for (
                            src_idx,
                            tgt_idx,
                            _,
                            _,
                            _,
                        ) in col_mappings:
                            cell_val = po_lookup[po_key][src_idx]
                            # Clean whitespace/trailing tabs if string
                            if isinstance(cell_val, str):
                                cell_val = cell_val.strip()
                            ws_b.cell(row=r, column=tgt_idx).value = cell_val
                        updated_count += 1
                    else:
                        missing_records.append(
                            {"Row in File B": r, "Customer PO": po_key}
                        )

                # 4. Save to in-memory bytes buffer
                output_buffer = io.BytesIO()
                wb_b.save(output_buffer)
                output_buffer.seek(0)

                # 5. Display success & statistics
                st.success(
                    f"✅ Successfully matched and updated **{updated_count}** rows in File B!"
                )

                metric_c1, metric_c2, metric_c3 = st.columns(3)
                metric_c1.metric("Indexed POs in File A", len(po_lookup))
                metric_c2.metric("Updated Rows in File B", updated_count)
                metric_c3.metric(
                    "Unmatched / Missing POs", len(missing_records)
                )

                if missing_records:
                    with st.expander(
                        f"⚠️ View {len(missing_records)} Unmatched PO(s)"
                    ):
                        st.table(pd.DataFrame(missing_records))

                # 6. Download Button
                original_name = file_b.name
                output_filename = (
                    original_name.replace(".xlsx", "_updated.xlsx")
                    if original_name.endswith(".xlsx")
                    else f"updated_{original_name}"
                )

                st.download_button(
                    label="📥 Download Updated File B",
                    data=output_buffer,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

            except Exception as e:
                st.error(f"❌ An error occurred during processing: {str(e)}")