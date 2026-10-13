import openpyxl
from openpyxl.utils import column_index_from_string


def update_excel_data(file_a_path, file_b_path, output_path):
    print("Loading File A (Lookup source)...")
    # Load File A with data_only=True to read evaluated formula values
    wb_a = openpyxl.load_workbook(file_a_path, data_only=True)

    # Use the 'Tag' sheet if present, otherwise default to the active sheet
    ws_a = wb_a["Tag"] if "Tag" in wb_a.sheetnames else wb_a.active

    # File A lookup mapping: Column B contains the Customer PO numbers
    # Extract columns E, F, G, H, I, O
    col_e = column_index_from_string("E")
    col_f = column_index_from_string("F")
    col_g = column_index_from_string("G")
    col_h = column_index_from_string("H")
    col_i = column_index_from_string("I")
    col_o = column_index_from_string("O")

    po_lookup = {}
    for row in range(2, ws_a.max_row + 1):
        po_val = ws_a.cell(row=row, column=2).value  # Column B
        if po_val is not None:
            po_key = str(po_val).strip()
            # Only store the first matching instance per PO
            if po_key not in po_lookup:
                po_lookup[po_key] = {
                    "E": ws_a.cell(row=row, column=col_e).value,
                    "F": ws_a.cell(row=row, column=col_f).value,
                    "G": ws_a.cell(row=row, column=col_g).value,
                    "H": ws_a.cell(row=row, column=col_h).value,
                    "I": ws_a.cell(row=row, column=col_i).value,
                    "O": ws_a.cell(row=row, column=col_o).value,
                }

    print(f"Indexed {len(po_lookup)} unique PO entries from File A.")

    print("Loading File B (Target file)...")
    # Load File B with data_only=False to preserve existing formulas and styles
    wb_b = openpyxl.load_workbook(file_b_path, data_only=False)
    ws_b = wb_b.active

    # Target columns in File B:
    # A(E) -> B(S), A(F) -> B(T), A(G) -> B(L),
    # A(H) -> B(M), A(I) -> B(K), A(O) -> B(R)
    target_mapping = {
        "E": column_index_from_string("S"),
        "F": column_index_from_string("T"),
        "G": column_index_from_string("L"),
        "H": column_index_from_string("M"),
        "I": column_index_from_string("K"),
        "O": column_index_from_string("R"),
    }

    col_po_b = column_index_from_string("I")  # Column I in File B
    updated_rows = 0
    missing_pos = []

    # Iterate starting from row 25 in File B
    for row in range(25, ws_b.max_row + 1):
        po_cell_val = ws_b.cell(row=row, column=col_po_b).value

        # Stop or skip if the cell is empty
        if po_cell_val is None or not str(po_cell_val).strip():
            continue

        po_key = str(po_cell_val).strip()

        if po_key in po_lookup:
            data = po_lookup[po_key]
            for source_col, target_col in target_mapping.items():
                # Clean trailing tab/whitespace characters from string values
                val = data[source_col]
                if isinstance(val, str):
                    val = val.strip()
                ws_b.cell(row=row, column=target_col).value = val
            updated_rows += 1
        else:
            missing_pos.append((row, po_key))

    print(f"Successfully updated {updated_rows} rows in File B.")
    if missing_pos:
        print(f"Warning: {len(missing_pos)} POs were not found in File A:")
        for r, po in missing_pos:
            print(f"  Row {r}: {po}")

    print(f"Saving output to: {output_path}")
    wb_b.save(output_path)
    print("Done!")


if __name__ == "__main__":
    file_a = "MHJ SPO SHIPMENT TAG LABLE - 10014.xlsx"
    file_b = "MJ-26001560.xlsx"
    output_file = "MJ-26001560_OUTPUT.xlsx"

    update_excel_data(file_a, file_b, output_file)